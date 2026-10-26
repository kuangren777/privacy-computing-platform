#!/usr/bin/env python3
"""验收进度表生成器。

把三类信息合成一张进度表：
  1. docs/acceptance/items.json           验收主表（合同 C/T 项 + V5 三级功能 + 合同补充项，固定编号）
  2. docs/acceptance/<模块>.json           各模块开发者填的实现映射（页面、操作步骤、接口、测试、工程实现说明、status）
  3. 测试结果                              node --test 的 TAP 输出（测试名以 [编号] 开头）+ 浏览器验收结果 e2e-results.json

当前报告仅核对前端证据：实现映射 + 该编号单元测试 + 同源码真实后端浏览器记录。
旧Mock单元/浏览器结果不作为当前证据；前端通过不证明后端算法或合同指标验收。

用法：
  python3 tools/acceptance.py            # 跑测试并生成 dist/acceptance/{index.html,status.json,验收进度表.xlsx}
  python3 tools/acceptance.py --no-test  # 复用上次的 docs/acceptance/unit-results.json
"""
import argparse
import datetime as dt
import html
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ACC = ROOT / 'docs' / 'acceptance'
OUT = ROOT / 'dist' / 'acceptance'
SGT = dt.timezone(dt.timedelta(hours=8))
CONSOLE_URL = 'http://127.0.0.1:8412/'
UNIT_SCOPE = 'real-client-isolated-unit'
BROWSER_SCOPE = 'real-backend-browser'
RESULT_SCHEMA = 'frontend-unit-evidence/1'

MODULE_ORDER = ['FL', 'HE', 'MPC', 'PSI', 'DP']
MODULE_CN = {'FL': '联邦学习', 'HE': '同态加密', 'MPC': '多方安全计算', 'PSI': '隐私求交', 'DP': '差分隐私'}
ID_RX = re.compile(r'\[([A-Z]+-[A-Z]+-\d{2})\]')


def source_digest():
    """Bind evidence to executable client/proxy bytes and the tests actually run."""
    names = [path for path in (ROOT / 'web').rglob('*')
             if path.is_file() and path.suffix in ('.js', '.mjs', '.css', '.html', '.json', '.svg')
             and not any(part in ('node_modules', '__pycache__') for part in path.parts)]
    names.extend(path for path in (ROOT / 'tests').glob('test_*.py') if path.is_file())
    names.extend(path for path in (ROOT / 'scripts').glob('*.py') if path.is_file())
    names.append(ROOT / 'tools' / 'acceptance.py')
    fingerprints = {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest()
                    for path in sorted(set(names)) if path.is_file()}
    return hashlib.sha256(json.dumps(fingerprints, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def parse_tap(stdout, exit_code):
    by_id, cases = {}, []
    for line in stdout.splitlines():
        match = re.match(r'^\s*(not ok|ok) \d+ - (.*)$', line)
        if not match:
            continue
        name = match[2].strip()
        directive = re.search(r'\s+#\s+(SKIP|TODO)\b', name, re.I)
        skipped, todo = bool(directive and directive[1].upper() == 'SKIP'), bool(directive and directive[1].upper() == 'TODO')
        if directive:
            name = name[:directive.start()].strip()
        passed = match[1] == 'ok' and not skipped and not todo
        case = {'name': name, 'pass': passed, 'skipped': skipped, 'todo': todo,
                'failed': match[1] == 'not ok' and not skipped and not todo}
        cases.append(case)
        for identity in ID_RX.findall(name):
            by_id.setdefault(identity, []).append(case)
    if not cases:
        raise ValueError('Node没有产生测试用例记录，不能生成通过结果')
    return {'schema': RESULT_SCHEMA, 'scope': UNIT_SCOPE,
            'ran_at': dt.datetime.now(SGT).isoformat(timespec='seconds'),
            'total': len(cases), 'passed': sum(case['pass'] for case in cases),
            'failed': sum(case['failed'] for case in cases), 'skipped': sum(case['skipped'] for case in cases),
            'todo': sum(case['todo'] for case in cases), 'exit_code': exit_code,
            'by_id': by_id, 'unmapped': [case for case in cases if not ID_RX.findall(case['name'])],
            'limitation': 'Isolated client transport/DOM tests; not native backend execution, browser acceptance or 129-route coverage.'}


def run_unit_tests(destination=None, test_log=None):
    tests = sorted(str(path.relative_to(ROOT / 'web')) for path in (ROOT / 'web' / 'tests').glob('*.test.mjs'))
    if not tests:
        raise ValueError('web/tests下没有当前客户端测试文件')
    digest = source_digest()
    command = ['node', '--test', '--test-reporter=tap'] + tests
    proc = subprocess.run(command, cwd=ROOT / 'web', capture_output=True, text=True)
    if test_log:
        test_log = Path(test_log); test_log.parent.mkdir(parents=True, exist_ok=True)
        test_log.write_text(proc.stdout, encoding='utf-8')
        test_log.with_suffix(test_log.suffix+'.stderr').write_text(proc.stderr, encoding='utf-8')
    result = parse_tap(proc.stdout, proc.returncode)
    result.update(source_sha256=digest, command=command)
    if source_digest() != digest:
        result.update(exit_code=1, source_changed_during_tests=True)
    destination = Path(destination) if destination else ACC / 'unit-results.json'
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding='utf-8')
    return result


def current_unit(unit):
    return (isinstance(unit, dict) and unit.get('schema') == RESULT_SCHEMA
            and unit.get('scope') == UNIT_SCOPE and unit.get('source_sha256') == source_digest()
            and isinstance(unit.get('by_id'), dict)
            and isinstance(unit.get('unmapped'), list)
            and all(type(unit.get(key)) is int and unit[key] >= 0 for key in ('total','passed','failed','skipped','todo','exit_code'))
            and not unit.get('source_changed_during_tests'))


def current_browser(document):
    return (isinstance(document, dict) and document.get('scope') == BROWSER_SCOPE
            and document.get('source_sha256') == source_digest() and isinstance(document.get('by_id'), dict))


def load_json(p, default):
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else default


def build(unit, browser_path=None):
    master = load_json(ACC / 'items.json', None)
    impl = {}
    for p in sorted(ACC.glob('*.json')):
        if p.name in ('items.json', 'unit-results.json', 'e2e-results.json'):
            continue
        impl.update(load_json(p, {}))
    browser_document = load_json(Path(browser_path) if browser_path else ACC / 'e2e-results.json', {})
    e2e = browser_document.get('by_id', {}) if current_browser(browser_document) else {}
    valid_unit = current_unit(unit)

    rows = []
    for it in master['items']:
        im = impl.get(it['id'], {})
        ut = unit.get('by_id', {}).get(it['id'], []) if valid_unit else []
        ut_pass = unit.get('exit_code') == 0 and bool(ut) and all(x['pass'] and not x.get('skipped') and not x.get('todo') for x in ut)
        ev = e2e.get(it['id'])
        # partial = 功能已实现、数值为原型内置评估引擎仿真结果：即使前端测试与浏览器均通过，也不能据此确认后端合同指标
        if im.get('status') in ('done', 'partial') and ut_pass and ev and ev.get('pass') is True:
            state = 'done'
        elif ut and not ut_pass or (ev and not ev.get('pass')):
            state = 'fail'
        elif im.get('status') in ('done', 'partial') and ut_pass:
            state = 'testing'
        elif im:
            state = 'unverified'
        else:
            state = 'todo'
        rows.append(dict(it, impl=im, unit=ut, e2e=ev, state=state,
                         contract=[c['id'] for c in master['contract'] if it['id'] in c['maps']]))
    by = {r['id']: r for r in rows}
    contract = []
    for c in master['contract']:
        sub = [by[m] for m in c['maps']]
        done = sum(1 for s in sub if s['state'] == 'done')
        contract.append(dict(c, done=done, total=len(sub)))
    return rows, contract


STATE_CN = {'unverified': '缺当前映射证据', 'done': '前端已验证', 'testing': '待验收', 'dev': '开发中', 'fail': '测试未通过', 'todo': '未开始'}


def render_html(rows, contract, unit, generated):
    def e(s):
        return html.escape(str(s or ''))

    done = sum(r['state'] == 'done' for r in rows)
    v5 = [r for r in rows if r['source'] == 'V5']
    v5done = sum(r['state'] == 'done' for r in v5)
    cdone = sum(c['done'] == c['total'] for c in contract)

    mod_stats = []
    for m in MODULE_ORDER:
        rs = [r for r in rows if r['id'].startswith(m + '-')]
        mod_stats.append((m, sum(r['state'] == 'done' for r in rs), len(rs)))

    def check(state):
        if state == 'done':
            return '<span class="ck on" title="前端已验证"><svg viewBox="0 0 16 16"><path d="M3.5 8.4l3 3 6-6.4"/></svg></span>'
        return '<span class="ck" title="%s"></span>' % STATE_CN[state]

    crow = []
    for c in contract:
        pct = round(100 * c['done'] / c['total']) if c['total'] else 0
        links = ' '.join('<a href="#%s">%s</a>' % (m, m) for m in c['maps'])
        crow.append(
            '<tr><td class="mono">%s</td><td>%s</td><td><b>%s</b><div class="sub">%s</div></td>'
            '<td class="links">%s</td><td><div class="pbar"><i style="width:%d%%"></i></div>'
            '<span class="pnum">%d / %d 项</span></td></tr>' % (
                e(c['id']), e(c['module']), e(c['name']), e(c['desc']), links, pct, c['done'], c['total']))

    irow = []
    for r in rows:
        im = r['impl']
        steps = ''.join('<li>%s</li>' % e(s) for s in im.get('steps', []))
        apis = ''.join('<code>%s</code>' % e(a) for a in im.get('api', []))
        tests = ''.join('<li class="%s">%s</li>' % ('t-ok' if t['pass'] else 't-bad', e(t['name'])) for t in r['unit'])
        ev = r['e2e']
        evh = ''
        if ev:
            evh = '<p><b>浏览器验收：</b>%s（%s SGT）%s</p>' % (
                '通过' if ev.get('pass') else '未通过', e(ev.get('checked_at', '')),
                ' · %s' % e(ev['note']) if ev.get('note') else '')
        detail = (
            '<div class="detail"><div class="dgrid">'
            '<div><h4>甲方要求（%s）</h4><p>%s</p>%s</div>'
            '<div><h4>工程实现说明</h4><p>%s</p></div>'
            '<div><h4>操作路径</h4><p>%s</p><ol>%s</ol><p><b>预期结果：</b>%s</p></div>'
            '<div><h4>接口</h4><p class="apis">%s</p><h4>测试用例</h4><ul class="tests">%s</ul>%s</div>'
            '</div></div>') % (
            '工程化清单 V5' if r['source'] == 'V5' else '合同补充项', e(r['desc']),
            ('<p class="sub">对应合同条目：%s</p>' % e('、'.join(r['contract']))) if r['contract'] else '<p class="sub">不对应合同功能条目（V5 工程化补充）</p>',
            e(im.get('impl') or '（开发中）'),
            e(im.get('ui') or ''), steps, e(im.get('expected') or ''),
            apis or '—', tests or '<li>暂无</li>', evh)
        irow.append(
            '<tbody class="item" id="%s" data-mod="%s" data-state="%s"><tr class="row">'
            '<td>%s</td><td class="mono">%s</td><td>%s</td><td>%s</td><td><b>%s</b>%s</td>'
            '<td><span class="st st-%s">%s</span></td><td class="mono">%d / %d</td>'
            '<td><a href="%s%s" target="_blank" rel="noopener">打开</a></td></tr>'
            '<tr class="drow"><td colspan="8">%s</td></tr></tbody>' % (
                e(r['id']), r['id'].split('-')[0], r['state'], check(r['state']), e(r['id']),
                e(r['l1']), e(r['l2']), e(r['l3']),
                ('' if r['source'] == 'V5' else ' <span class="tag">合同补充</span>') +
                (' <span class="tag">仿真数值</span>' if im.get('status') == 'partial' else ''),
                r['state'], STATE_CN[r['state']], sum(t['pass'] for t in r['unit']), len(r['unit']),
                CONSOLE_URL, e(r['route']), detail))

    mods = ''.join(
        '<div class="mod"><span>%s</span><b>%d / %d</b><div class="pbar"><i style="width:%d%%"></i></div></div>' % (
            MODULE_CN[m], d, t, round(100 * d / t) if t else 0) for m, d, t in mod_stats)
    filters = ''.join('<button type="button" data-f="%s">%s</button>' % (m, MODULE_CN[m]) for m in MODULE_ORDER)

    return TEMPLATE.format(
        generated=e(generated), done=done, total=len(rows), v5done=v5done, v5=len(v5),
        cdone=cdone, ctotal=len(contract), utotal=unit['total'], ufail=unit['failed'],
        uran=e(unit['ran_at'].replace('T', ' ')[:16]), mods=mods, filters=filters,
        crows=''.join(crow), irows=''.join(irow), console=CONSOLE_URL)


TEMPLATE = '''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>验收进度表</title>
<style>
:root {{ --bg:#F5F6F8; --card:#fff; --ink:#14171C; --muted:#6E7787; --line:#E1E5EB; --row:#F3F5F8;
  --primary:#1D4E89; --soft:#EDF2F9; --ok:#1D5C42; --okbg:#E9F4EF; --warn:#9A6516; --warnbg:#FBF3E4;
  --bad:#A33B21; --badbg:#FBEDEA; --idle:#4C5564; --idlebg:#F0F2F5; }}
@media (prefers-color-scheme: dark) {{ :root:not([data-theme="light"]) {{ --bg:#111418; --card:#191D23; --ink:#E6E9EE; --muted:#98A1AF;
  --line:#2A3039; --row:#20252D; --primary:#7FA9DB; --soft:#1E2A3A; --ok:#6CC59A; --okbg:#173327; --warn:#E0B060; --warnbg:#3A2E17;
  --bad:#F08A70; --badbg:#3B1F19; --idle:#B4BCC8; --idlebg:#262B33; }} }}
:root[data-theme="dark"] {{ --bg:#111418; --card:#191D23; --ink:#E6E9EE; --muted:#98A1AF; --line:#2A3039; --row:#20252D; --primary:#7FA9DB; --soft:#1E2A3A; --ok:#6CC59A; --okbg:#173327; --warn:#E0B060; --warnbg:#3A2E17; --bad:#F08A70; --badbg:#3B1F19; --idle:#B4BCC8; --idlebg:#262B33; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--bg); color:var(--ink); font:14px/1.6 "PingFang SC","Microsoft YaHei",system-ui,sans-serif; }}
.wrap {{ max-width:1280px; margin:0 auto; padding:28px 16px 60px; }}
h1 {{ font-size:22px; margin:0 0 4px; }} h2 {{ font-size:16px; margin:28px 0 10px; }}
.sub {{ color:var(--muted); font-size:12.5px; }}
.kpis {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(200px,1fr)); gap:12px; margin-top:18px; }}
.kpi {{ background:var(--card); border:1px solid var(--line); border-radius:8px; padding:14px 16px; }}
.kpi b {{ font-size:24px; display:block; }} .kpi span {{ color:var(--muted); font-size:12.5px; }}
.mods {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(180px,1fr)); gap:10px; margin-top:12px; }}
.mod {{ background:var(--card); border:1px solid var(--line); border-radius:8px; padding:10px 14px; display:grid; grid-template-columns:1fr auto; gap:6px; }}
.mod .pbar {{ grid-column:1/3; }}
.pbar {{ height:6px; background:var(--row); border-radius:3px; overflow:hidden; min-width:80px; }}
.pbar i {{ display:block; height:100%; background:var(--ok); }}
.pnum {{ font-size:12px; color:var(--muted); }}
.card {{ background:var(--card); border:1px solid var(--line); border-radius:8px; overflow-x:auto; }}
table {{ border-collapse:collapse; width:100%; min-width:860px; }}
th, td {{ text-align:left; padding:9px 12px; border-bottom:1px solid var(--row); vertical-align:top; }}
th {{ font-size:12px; color:var(--muted); font-weight:600; background:var(--card); position:sticky; top:0; }}
.mono {{ font-family:ui-monospace,Menlo,Consolas,monospace; font-size:12.5px; white-space:nowrap; }}
.links a {{ margin-right:6px; font-size:12px; white-space:nowrap; }}
a {{ color:var(--primary); }}
.ck {{ width:18px; height:18px; border:1.6px solid var(--line); border-radius:4px; display:inline-block; vertical-align:middle; }}
.ck.on {{ background:var(--ok); border-color:var(--ok); }}
.ck svg {{ width:14px; height:14px; margin:1px; fill:none; stroke:#fff; stroke-width:2.2; }}
.st {{ font-size:12px; padding:2px 8px; border-radius:4px; white-space:nowrap; }}
.st-done {{ background:var(--okbg); color:var(--ok); }} .st-testing {{ background:var(--soft); color:var(--primary); }}
.st-unverified {{ background:var(--idlebg); color:var(--idle); }} .st-dev {{ background:var(--warnbg); color:var(--warn); }} .st-fail {{ background:var(--badbg); color:var(--bad); }}
.st-todo {{ background:var(--idlebg); color:var(--idle); }}
.tag {{ font-size:11px; padding:1px 6px; border-radius:3px; background:var(--soft); color:var(--primary); margin-left:6px; font-weight:400; }}
.row {{ cursor:pointer; }} .row:hover td {{ background:var(--row); }}
.drow {{ display:none; }} .item.open .drow {{ display:table-row; }}
.detail {{ padding:4px 4px 10px; }}
.dgrid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(260px,1fr)); gap:14px 22px; }}
.detail h4 {{ margin:6px 0 4px; font-size:12.5px; color:var(--muted); }} .detail p, .detail ol, .detail ul {{ margin:0 0 6px; }}
.detail ol, .detail ul {{ padding-left:18px; }}
.apis code {{ display:inline-block; margin:0 6px 4px 0; padding:1px 6px; background:var(--row); border-radius:3px; font-size:12px; }}
.tests li {{ font-size:12.5px; }} .t-ok {{ color:var(--ok); }} .t-bad {{ color:var(--bad); }}
.bar {{ display:flex; gap:8px; flex-wrap:wrap; align-items:center; margin:10px 0; }}
.bar button {{ border:1px solid var(--line); background:var(--card); color:var(--ink); border-radius:5px; padding:4px 10px; font:inherit; font-size:13px; cursor:pointer; }}
.bar button.on {{ border-color:var(--primary); color:var(--primary); background:var(--soft); }}
.bar input {{ border:1px solid var(--line); background:var(--card); color:var(--ink); border-radius:5px; padding:5px 10px; font:inherit; font-size:13px; min-width:200px; }}
.legend {{ font-size:12.5px; color:var(--muted); margin-top:8px; }}
</style></head><body><div class="wrap">
<h1>隐私计算联邦学习软件原型 · 验收进度表</h1>
<div class="sub">2026 年 AI+稀土追溯模型关键技术及应用研究（隐私计算）联合研发项目 · 生成于 {generated} SGT · 真实客户端入口：<a href="{console}" target="_blank" rel="noopener">控制台</a> · <a href="验收进度表.xlsx">下载 Excel 版</a></div>
<div class="kpis">
 <div class="kpi"><b>{done} / {total}</b><span>具备当前前端映射证据（不是后端合同验收）</span></div>
 <div class="kpi"><b>{v5done} / {v5}</b><span>V5 前端关联项已验证</span></div>
 <div class="kpi"><b>{cdone} / {ctotal}</b><span>合同关联前端项已验证（不代表指标达标）</span></div>
 <div class="kpi"><b>{utotal} 个 · 失败 {ufail} 个</b><span>自动化测试用例（最近一次运行 {uran} SGT）</span></div>
</div>
<div class="mods">{mods}</div>
<div class="legend">「前端已验证」仅指当前源码下带编号的客户端测试与真实后端浏览器记录匹配；无编号测试不自动分配到验收项。旧Mock证据不复用，后端原生算法、129个接口和合同指标需单独验收。标「仿真数值」的历史映射仅说明原有评估口径，不能据此判断当前原生算法或指标达标。</div>

<h2>一、合同条目对照（附件 1 功能清单 + 技术参数指标）</h2>
<div class="card"><table><thead><tr><th>编号</th><th>一级模块</th><th>二级模块 / 合同说明</th><th>对应三级功能</th><th>完成度</th></tr></thead><tbody>{crows}</tbody></table></div>

<h2>二、三级功能逐项进度（点击一行展开：工程实现说明、操作路径、接口、测试用例）</h2>
<div class="bar"><button type="button" data-f="" class="on">全部</button>{filters}
 <span style="flex:1"></span><button type="button" data-s="pending">只看未完成</button><input id="q" type="search" placeholder="搜索编号或功能名"></div>
<div class="card"><table><thead><tr><th style="width:44px">完成</th><th>编号</th><th>一级</th><th>二级</th><th>三级功能</th><th>状态</th><th>测试通过</th><th>演示</th></tr></thead>{irows}</table></div>
</div>
<script>
(function () {{
  var mod = '', pending = false, q = '';
  function apply() {{
    document.querySelectorAll('tbody.item').forEach(function (b) {{
      var ok = (!mod || b.dataset.mod === mod) && (!pending || b.dataset.state !== 'done') &&
        (!q || b.textContent.toLowerCase().indexOf(q) >= 0);
      b.hidden = !ok;
    }});
  }}
  document.querySelectorAll('[data-f]').forEach(function (b) {{
    b.addEventListener('click', function () {{
      mod = b.dataset.f;
      document.querySelectorAll('[data-f]').forEach(function (x) {{ x.classList.toggle('on', x === b); }});
      apply();
    }});
  }});
  var pb = document.querySelector('[data-s]');
  pb.addEventListener('click', function () {{ pending = !pending; pb.classList.toggle('on', pending); apply(); }});
  document.getElementById('q').addEventListener('input', function (e) {{ q = e.target.value.trim().toLowerCase(); apply(); }});
  document.querySelectorAll('tr.row').forEach(function (r) {{
    r.addEventListener('click', function (ev) {{ if (ev.target.tagName !== 'A') r.parentNode.classList.toggle('open'); }});
  }});
  if (location.hash) {{ var t = document.getElementById(location.hash.slice(1)); if (t) t.classList.add('open'); }}
}})();
</script></body></html>
'''


def write_xlsx(rows, contract, path):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print('openpyxl 未安装，跳过 xlsx', file=sys.stderr)
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '三级功能进度'
    head = ['完成', '编号', '一级功能', '二级功能', '三级功能', '来源', '甲方要求', '对应合同条目', '状态',
            '操作路径', '接口', '测试用例（通过/总数）', '工程实现说明']
    ws.append(head)
    for r in rows:
        im = r['impl']
        ws.append(['☑' if r['state'] == 'done' else '☐', r['id'], r['l1'], r['l2'], r['l3'], r['source'], r['desc'],
                   '、'.join(r['contract']), STATE_CN[r['state']], im.get('ui', ''), '\n'.join(im.get('api', [])),
                   '%d / %d' % (sum(t['pass'] for t in r['unit']), len(r['unit'])), im.get('impl', '')])
    ws2 = wb.create_sheet('合同条目对照')
    ws2.append(['编号', '一级模块', '二级模块', '合同说明', '对应三级功能', '前端已验证 / 关联项'])
    for c in contract:
        ws2.append([c['id'], c['module'], c['name'], c['desc'], '、'.join(c['maps']), '%d / %d' % (c['done'], c['total'])])
    font = Font(name='宋体', size=10.5)
    bold = Font(name='宋体', size=10.5, bold=True)
    for w in (ws, ws2):
        for row in w.iter_rows():
            for cell in row:
                cell.font = bold if cell.row == 1 else font
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in w[1]:
            cell.fill = PatternFill('solid', fgColor='E1E5EB')
        w.freeze_panes = 'A2'
    for col, wd in zip('ABCDEFGHIJKLM', [6, 13, 12, 12, 18, 9, 40, 12, 10, 30, 30, 12, 50]):
        ws.column_dimensions[col].width = wd
    for col, wd in zip('ABCDEF', [8, 14, 22, 50, 40, 12]):
        ws2.column_dimensions[col].width = wd
    wb.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--no-test', action='store_true')
    parser.add_argument('--output', type=Path, default=OUT)
    parser.add_argument('--unit-results', type=Path, default=ACC/'unit-results.json')
    parser.add_argument('--e2e-results', type=Path, default=ACC/'e2e-results.json')
    parser.add_argument('--test-log', type=Path)
    args = parser.parse_args(argv)
    try:
        unit = load_json(args.unit_results, None) if args.no_test else run_unit_tests(args.unit_results, args.test_log)
        if not current_unit(unit):
            raise ValueError('缓存测试结果缺少当前源码/real-client scope，请不带--no-test重新运行')
        rows, contract = build(unit, args.e2e_results)
        args.output.mkdir(parents=True, exist_ok=True)
        generated = dt.datetime.now(SGT).strftime('%Y-%m-%d %H:%M')
        (args.output/'index.html').write_text(render_html(rows, contract, unit, generated), encoding='utf-8')
        browser = load_json(args.e2e_results, {})
        (args.output/'status.json').write_text(json.dumps({
            'schema': 'frontend-acceptance-progress/1', 'scope': 'frontend-evidence-only',
            'source_sha256': unit['source_sha256'], 'generated_at': generated+' SGT',
            'unit': {key: unit[key] for key in ('scope','total','passed','failed','skipped','todo','exit_code')},
            'unmapped_unit_cases': len(unit['unmapped']),
            'current_browser_evidence': current_browser(browser),
            'historical_browser_records_ignored': bool(browser.get('by_id')) and not current_browser(browser),
            'limitation': 'Client-only evidence. No native-backend, full-API or contract acceptance claim.',
            'items': [{key: row[key] for key in ('id','l1','l2','l3','source','state')} for row in rows],
            'contract': [{key: row[key] for key in ('id','name','done','total')} for row in contract]
        }, ensure_ascii=False, indent=1), encoding='utf-8')
        write_xlsx(rows, contract, args.output/'验收进度表.xlsx')
        print('前端单元测试 %d（失败%d，跳过%d，未映射编号%d）；当前浏览器证据=%s → %s' % (
            unit['total'], unit['failed'], unit['skipped'], len(unit['unmapped']), current_browser(browser), args.output))
        return 0 if unit['exit_code'] == 0 and not unit['failed'] else 1
    except (OSError, ValueError, KeyError, TypeError) as error:
        print('验收资料生成失败：'+str(error), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
